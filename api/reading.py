"""독서실 신청서 / 대기자 엑셀 생성 — 웹에서 읽은 명단(JSON)을 양식(_templates_reading/독서실.xlsx)에 채워 돌려줌."""
import base64, io, json, re, datetime as dt
from pathlib import Path
from http.server import BaseHTTPRequestHandler
import openpyxl

TPL = Path(__file__).resolve().parent / "_templates_reading" / "독서실.xlsx"
KINDS = ["개인", "남자일반", "여자일반"]


def _d(s):
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(s or ""))
    return dt.datetime(int(m[1]), int(m[2]), int(m[3])) if m else None


def retitle(ws, y, sm):
    em = sm + 1
    last = dt.date(y + em // 12, em % 12 + 1, 1) - dt.timedelta(days=1)
    yy = y % 100
    a1 = str(ws["A1"].value or "")
    ws["A1"] = re.sub(r"\d{1,2}\s*~\s*\d{1,2}월", f"{sm}~{em}월", a1)
    ws["A3"] = f"사용기간 :  {yy:02d}-{sm:02d}-01 ~ {last.year % 100:02d}-{em:02d}-{last.day:02d}"
    a5 = str(ws["A5"].value or "")
    ws["A5"] = re.sub(r"다음추첨:\s*\d{1,2}월", f"다음추첨: {em}월", a5)


def fill_apply(ws, rows, log, kind):
    seat_row = {ws.cell(r, 1).value: r for r in range(11, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int)}
    for p in sorted(rows, key=lambda p: int(p.get("seat") or 0)):
        r = seat_row.get(int(p.get("seat") or 0))
        if not r:
            log.append(f"⚠ {kind}: 좌석 {p.get('seat')}번이 양식에 없어 제외 — {p.get('name')}")
            continue
        if ws.cell(r, 5).value:
            log.append(f"⚠ {kind}: 좌석 {p['seat']}번 중복 — {p.get('name')} 제외")
            continue
        ws.cell(r, 2).value = _d(p.get("date"))
        ws.cell(r, 3).value, ws.cell(r, 4).value = p.get("dong"), p.get("ho")
        ws.cell(r, 5).value, ws.cell(r, 6).value = p.get("name"), p.get("phone")
    log.append(f"{kind} 신청서: {len(rows)}명")


def fill_wait(ws, rows, log, kind):
    slots = [r for r in range(11, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int)]
    rows = sorted(rows, key=lambda p: int(p.get("order") or 0))
    if len(rows) > len(slots):
        log.append(f"⚠ {kind}: 대기자 {len(rows)}명 중 칸이 {len(slots)}개라 뒤 {len(rows) - len(slots)}명 제외")
    for r, p in zip(slots, rows):
        ws.cell(r, 2).value = _d(p.get("date"))
        ws.cell(r, 3).value, ws.cell(r, 4).value = p.get("dong"), p.get("ho")
        ws.cell(r, 5).value, ws.cell(r, 6).value = p.get("name"), p.get("phone")
    log.append(f"{kind} 대기자: {min(len(rows), len(slots))}명")


def build(b):
    y, sm = int(b["year"]), int(b["start"])
    if not 1 <= sm <= 11:
        raise ValueError("시작 월이 올바르지 않아요.")
    wb = openpyxl.load_workbook(TPL)
    log = []
    for k in KINDS:
        for name in (k, f"{k}_대기"):
            retitle(wb[name], y, sm)
        d = (b.get("sheets") or {}).get(k) or {}
        if d.get("apply"):
            fill_apply(wb[k], d["apply"], log, k)
        if d.get("wait"):
            fill_wait(wb[f"{k}_대기"], d["wait"], log, k)
    buf = io.BytesIO()
    wb.save(buf)
    return {"file": base64.b64encode(buf.getvalue()).decode(), "log": log}


class handler(BaseHTTPRequestHandler):
    def _send(self, code, obj):
        raw = json.dumps(obj, ensure_ascii=False).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.end_headers()
        self.wfile.write(raw)

    def do_POST(self):
        try:
            n = int(self.headers.get("Content-Length") or 0)
            self._send(200, build(json.loads(self.rfile.read(n) or b"{}")))
        except Exception as e:  # noqa
            self._send(400, {"error": f"{type(e).__name__}: {e}"})
