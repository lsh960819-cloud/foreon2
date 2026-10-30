"""수영 반변경 — 신청서 엑셀 + 반변경 명단(양식) → 반변경이 반영된 신청서 엑셀.

POST /api/swapclass  {apply: <base64 xlsx 또는 신청서 zip>, apply_name, list: <base64 반변경 양식 xlsx>}
→ {file, file_name, log, results, over}

규칙
- 기존 반: 해당 회원(성함+동+호 모두 일치) 행을 지우고 아래 명단을 위로 당김 (날짜·서명·비고 등 같은 줄 내용도 함께 이동)
- 변경 반: 맨 밑 빈칸에 추가, 비고에 '반변경'
- 확실하지 않은 건 바꾸지 않고 사유를 log/results 에 남김 (추측 금지)
"""
import base64, io, json, re, zipfile, datetime as dt
from http.server import BaseHTTPRequestHandler
import openpyxl

LEVELS = ["기초초급", "중상급", "기초", "초급", "중급", "상급"]
DAYS = ["화목", "수금", "토"]
TITLE_RE = re.compile(r"(성인|어린이)수영-(기초초급|중상급|기초|초급|중급|상급)\((화목|수금|토)(\d{1,2})시(?:(\d{1,2})분)?(.*)\)")
CAP = 25


# ── 입력 정리 ──────────────────────────────────────────────
def _s(v):
    return re.sub(r"\s+", "", str(v if v is not None else ""))


def _num(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    d = re.sub(r"\D", "", str(v if v is not None else ""))
    return int(d) if d else None


def norm_time(v):
    """06:30 / 6:30 / 0630 / 6시30분 / 06시 / 9 / 엑셀 시간값 → (6, 30). 해석 불가면 None"""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.hour, v.minute
    if isinstance(v, dt.time):
        return v.hour, v.minute
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 0 < v < 1:                      # 엑셀 시간 일련값
            mins = round(v * 24 * 60)
            return mins // 60, mins % 60
        v = str(int(v)) if float(v).is_integer() else str(v)
    t = _s(v)
    m = re.fullmatch(r"(\d{1,2})시(?:(\d{1,2})분?)?", t) or re.fullmatch(r"(\d{1,2})[:.](\d{2})", t)
    if m:
        return int(m.group(1)), int(m.group(2) or 0)
    if re.fullmatch(r"\d{3,4}", t):
        return int(t[:-2]), int(t[-2:])
    if re.fullmatch(r"\d{1,2}", t):
        return int(t), 0
    return None


def norm_day(v):
    t = _s(v).replace(",", "").replace("/", "").replace("·", "").replace("요일", "")
    t = {"화": "화목", "목": "화목", "수": "수금", "금": "수금", "토": "토", "토요일": "토"}.get(t, t)
    return t if t in DAYS else None


def norm_level(v):
    t = _s(v).replace("반", "")
    return t if t in LEVELS else None


def tlabel(hm):
    return f"{hm[0]:02d}:{hm[1]:02d}" if hm else "?"


def cls_label(k):
    g, lv, day, hm = k
    return f"{g} {day} {tlabel(hm)} {lv}"


# ── 신청서 시트 ────────────────────────────────────────────
def swim_sheets(wb):
    out = {}
    for ws in wb.worksheets:
        m = TITLE_RE.search(ws.title)
        if m:
            key = (m[1], m[2], m[3], (int(m[4]), int(m[5] or 0)))
            out[key] = ws
    return out


def slots(ws):
    return [r for r in range(15, ws.max_row + 1) if str(ws.cell(r, 2).value or "").strip().isdigit()]


def records(ws):
    """이름이 있는 행만 (C~J 8칸)"""
    return [[ws.cell(r, c).value for c in range(3, 11)] for r in slots(ws) if _s(ws.cell(r, 4).value)]


def write(ws, recs):
    rs = slots(ws)
    for i, r in enumerate(rs):
        row = recs[i] if i < len(recs) else [None] * 8
        for j, v in enumerate(row):
            ws.cell(r, 3 + j).value = v


def same(rec, name, dong, ho):
    return _s(rec[1]) == name and _num(rec[2]) == dong and _num(rec[3]) == ho


def instructor(ws):
    m = TITLE_RE.search(ws.title)
    return m[6] if m else ""


# ── 반변경 양식 읽기 ───────────────────────────────────────
HEAD = {"성함": "name", "동": "dong", "호": "ho", "호수": "ho", "구분": "group",
        "현재요일": "fday", "현재시간": "ftime", "현재반": "flv",
        "변경요일": "tday", "변경시간": "ttime", "변경반": "tlv"}


def read_list(raw):
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb["반변경명단"] if "반변경명단" in wb.sheetnames else wb.worksheets[0]
    hr, cols = None, {}
    for r in range(1, min(ws.max_row, 20) + 1):
        row = {_s(ws.cell(r, c).value).replace("(", "").replace(")", ""): c for c in range(1, ws.max_column + 1)}
        hit = {HEAD[k]: c for k, c in row.items() if k in HEAD}
        if {"name", "flv", "tlv"} <= hit.keys():
            hr, cols = r, hit
            break
    if not hr:
        raise ValueError("반변경 명단에서 머리글(성함·현재 반·변경 반)을 못 찾았어요. 양식 파일을 써 주세요.")
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        g = {k: ws.cell(r, c).value for k, c in cols.items()}
        if not _s(g.get("name")):
            continue
        g["row"] = r
        rows.append(g)
    return rows


# ── 적용 ──────────────────────────────────────────────────
def apply_changes(wb, rows):
    sheets = swim_sheets(wb)
    if not sheets:
        raise ValueError("신청서 파일에서 수영 시트를 못 찾았어요. (시트 이름 예: 2단지 성인수영-초급(화목06시30분이규상))")
    log, results = [], []

    def where(name, dong, ho):
        return [cls_label(k) for k, ws in sheets.items() if any(same(x, name, dong, ho) for x in records(ws))]

    def options(g, day):
        return ", ".join(sorted({f"{tlabel(k[3])} {k[1]}" for k in sheets if k[0] == g and k[2] == day})) or "없음"

    for g in rows:
        name, dong, ho = _s(g.get("name")), _num(g.get("dong")), _num(g.get("ho"))
        grp = _s(g.get("group")) or "성인"
        fday, tday = norm_day(g.get("fday")), norm_day(g.get("tday")) or norm_day(g.get("fday"))
        ft = norm_time(g.get("ftime"))
        tt = norm_time(g.get("ttime")) or ft
        flv, tlv = norm_level(g.get("flv")), norm_level(g.get("tlv"))
        res = {"row": g["row"], "name": name, "dong": dong, "ho": ho,
               "from": f"{fday or '?'} {tlabel(ft)} {flv or '?'}", "to": f"{tday or '?'} {tlabel(tt)} {tlv or '?'}",
               "ok": False, "msg": ""}
        results.append(res)

        def fail(msg):
            res["msg"] = msg
            log.append(f"⚠ {g['row']}행 {name}: {msg}")

        if grp not in ("성인", "어린이"):
            fail(f"구분 '{g.get('group')}' → 성인/어린이 중 하나로 적어 주세요"); continue
        if dong is None or ho is None:
            fail("동·호수가 비어 있어요"); continue
        if not (fday and ft and flv):
            fail("현재 요일·시간·반 중 비었거나 알아볼 수 없는 칸이 있어요"); continue
        if not tlv:
            fail("변경 반이 비었거나 알아볼 수 없어요"); continue
        fk, tk = (grp, flv, fday, ft), (grp, tlv, tday, tt)
        if fk == tk:
            fail("현재 반과 변경 반이 같아요"); continue
        fs = sheets.get(fk)
        if not fs:
            fail(f"현재 반({cls_label(fk)})이 신청서에 없어요 — {grp} {fday} 개설반: {options(grp, fday)}"); continue
        ts = sheets.get(tk)
        if not ts:
            fail(f"변경 반({cls_label(tk)})이 신청서에 없어요 — {grp} {tday} 개설반: {options(grp, tday)}"); continue

        a = records(fs)
        idx = [i for i, x in enumerate(a) if same(x, name, dong, ho)]
        if not idx:
            at = where(name, dong, ho)
            near = [_s(x[1]) for x in a if _num(x[2]) == dong and _num(x[3]) == ho]
            if at:
                fail(f"현재 반에 없어요 — 실제로는 [{' / '.join(at)}]에 있어요. 현재 반을 고쳐 주세요")
            elif near:
                fail(f"현재 반에 같은 동·호수({dong}-{ho})로 '{', '.join(near)}' 님이 있어요. 이름을 확인해 주세요")
            else:
                fail("신청서 어디에도 같은 성함·동·호수가 없어요")
            continue
        if len(idx) > 1:
            fail("현재 반에 같은 성함·동·호수가 2명 이상이라 바꾸지 않았어요"); continue
        b = records(ts)
        if any(same(x, name, dong, ho) for x in b):
            fail(f"변경 반({cls_label(tk)})에 이미 있어요"); continue
        if len(b) >= len(slots(ts)):
            fail(f"변경 반({cls_label(tk)}) 칸이 다 찼어요 ({len(slots(ts))}칸)"); continue

        rec = a.pop(idx[0])
        write(fs, a)
        old = _s(rec[7])
        rec[7] = "반변경" if not old or old == "반변경" else f"{rec[7]}, 반변경"
        b.append(rec)
        write(ts, b)
        res["ok"] = True
        res["from"] = f"{cls_label(fk)}({instructor(fs)})"
        res["to"] = f"{cls_label(tk)}({instructor(ts)})"
        res["msg"] = f"{len(b)}번째로 추가"

    over = []
    for k, ws in sheets.items():
        n = len(records(ws))
        if n > CAP:
            over.append({"cls": f"{cls_label(k)}({instructor(ws)})", "n": n})
    order = {d: i for i, d in enumerate(DAYS)}
    over.sort(key=lambda o: (order.get(o["cls"].split()[1], 9), o["cls"]))
    ok = sum(r["ok"] for r in results)
    log.insert(0, f"반변경 {len(results)}명 중 {ok}명 적용 · {len(results) - ok}명 미적용")
    return log, results, over


def open_apply(raw, name):
    """신청서 xlsx 또는 신청서 생성 zip(안의 수영 파일)"""
    zf = zipfile.ZipFile(io.BytesIO(raw))
    names = zf.namelist()
    if "[Content_Types].xml" in names:
        return openpyxl.load_workbook(io.BytesIO(raw)), name
    cand = [n for n in names if n.lower().endswith(".xlsx") and "수영" in n] or [n for n in names if n.lower().endswith(".xlsx")]
    if not cand:
        raise ValueError("zip 안에 엑셀 파일이 없어요.")
    return openpyxl.load_workbook(io.BytesIO(zf.read(cand[0]))), cand[0].split("/")[-1]


def build(b):
    raw = base64.b64decode(b["apply"])
    wb, fname = open_apply(raw, b.get("apply_name") or "수영 신청서.xlsx")
    rows = read_list(base64.b64decode(b["list"]))
    if not rows:
        raise ValueError("반변경 명단에 적힌 사람이 없어요.")
    log, results, over = apply_changes(wb, rows)
    buf = io.BytesIO()
    wb.save(buf)
    stem = re.sub(r"\.xlsx$", "", fname, flags=re.I)
    return {"file": base64.b64encode(buf.getvalue()).decode(), "file_name": f"{stem}_반변경.xlsx",
            "log": log, "results": results, "over": over}


class handler(BaseHTTPRequestHandler):
    def _send(self, code, obj):
        raw = json.dumps(obj, ensure_ascii=False).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.end_headers()
        self.wfile.write(raw)

    def do_POST(self):
        try:
            n = int(self.headers.get("Content-Length") or 0)
            self._send(200, build(json.loads(self.rfile.read(n) or b"{}")))
        except Exception as e:  # noqa
            self._send(400, {"error": f"{type(e).__name__}: {e}"})
